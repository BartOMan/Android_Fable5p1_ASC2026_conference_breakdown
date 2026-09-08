#!/usr/bin/env python3
"""Build the normalized ASC 2026 dataset from the scraped planner tables + website spreadsheets."""
import json,re,html,collections,unicodedata,os,sys,base64,io,datetime
S=os.environ.get('ASC_SCRATCH','/tmp/claude-0/-home-user-Android-Fable5p1-ASC2026-conference-breakdown/22e803a5-16ae-5ba4-91b7-cbd9a897b501/scratchpad')
D=S+'/data/'; W=S+'/web/'
import openpyxl
def L(f): return [x for x in json.load(open(D+f)) if isinstance(x,dict)]
ag=L('agenda.json'); md=L('media.json'); sp=L('speakers.json')

def strip_tags(s):
    s=re.sub(r'<br\s*/?>','\n',s); s=re.sub(r'</p>','\n',s); s=re.sub(r'<[^>]+>','',s)
    s=html.unescape(s); s=re.sub(r'[ \t]+',' ',s); s=re.sub(r'\n\s*\n+','\n',s).strip()
    return s
def norm(s):
    s=unicodedata.normalize('NFKD',s); s=''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9 ]','',s.lower()).strip()
def name_key(n):
    n=re.sub(r'\*','',n); n=re.sub(r'\(.*?\)','',n)
    toks=[t for t in norm(n).split() if t]
    toks=[t for t in toks if t not in ('jr','sr','ii','iii','dr','prof')]
    if not toks: return ''
    if len(toks)==1: return toks[0]
    return toks[0]+' '+toks[-1]
COUNTRY_MAP={'Korea (the Republic of)':'South Korea','United Republic of':'Tanzania','Russian Federation':'Russia','Czechia':'Czech Republic','United States':'United States','USA':'United States','U.S.A.':'United States','UK':'United Kingdom','P.R. China':'China','Taiwan':'Taiwan'}
ORG_SUFFIX=re.compile(r'^(Inc\.?|Ltd\.?|LLC|GmbH|Co\.?|Corp\.?|S\.?A\.?|SpA|AG|B\.?V\.?|LTD|INC|Limited|Corporation|Company|LLP|S\.?r\.?l\.?)$',re.I)
DEPT_RE=re.compile(r'^(Department|Dept\.?|Faculty|Division|Graduate School|School of|Laboratory of|Lab\.? of|Institute for|Center for|Centre for|Section|Group|Unit|Physics|Engineering|Chemistry|R&D|Research Center|Research Centre|Cluster|College of|Program|Programme|Applied Physics|Materials Science|Electrical Engineering|Mechanical Engineering)\b',re.I)
ORG_RE=re.compile(r'\b(University|Universit|Institut|Laborator|\bLab\b|Center|Centre|College|School|Academy|Corporation|Inc\.|Ltd|GmbH|Co\.|Company|Agency|Research|Technolog|Organization|Foundation|Hospital|Sciences|Bureau|Association|Consortium|Council|Facility|CERN|KEK|NIST|CEA|INFN|CNRS|Fermilab|Politecnico|Polytechnic|Hochschule|Universidad|Università|Université|Universiteit|Universität|Group|Division|Department|Faculty)',re.I)
US_STATES=set('Alabama Alaska Arizona Arkansas California Colorado Connecticut Delaware Florida Georgia Hawaii Idaho Illinois Indiana Iowa Kansas Kentucky Louisiana Maine Maryland Massachusetts Michigan Minnesota Mississippi Missouri Montana Nebraska Nevada Ohio Oklahoma Oregon Pennsylvania Tennessee Texas Utah Vermont Virginia Washington Wisconsin Wyoming'.split())|{'New York','New Jersey','New Mexico','New Hampshire','North Carolina','North Dakota','South Carolina','South Dakota','West Virginia','Rhode Island','District of Columbia','DC','D.C.'}
REGIONS=set(US_STATES)
def _collect_regions():
    for m in md:
        d=m['description']; i=d.find('<hr>'); body=d[i+4:]
        mm=re.search(r'<p><i>(.*?)</i>\s*<br />(.*?)</p>',body,re.S)
        if not mm: continue
        for sup,txt in re.findall(r'<sup>(\d+)</sup>([^<]+)',mm.group(2)):
            parts=[p.strip() for p in html.unescape(txt).strip().rstrip(';').split(',') if p.strip()]
            parts2=[]
            for p in parts:
                if parts2 and ORG_SUFFIX.match(p): parts2[-1]+=', '+p
                else: parts2.append(p)
            if len(parts2)>=5: REGIONS.add(parts2[-2])
_collect_regions()
def parse_aff(t):
    t=html.unescape(t).strip().rstrip(';').strip()
    t=re.sub(r'\s+',' ',t)
    raw=[p.strip() for p in t.split(',') if p.strip()]
    parts=[]
    for p in raw:
        if parts and ORG_SUFFIX.match(p): parts[-1]=parts[-1]+', '+p
        else: parts.append(p)
    country=parts[-1] if parts else ''
    country=COUNTRY_MAP.get(country,country)
    region=''; city=''; inst=''; dept=''
    if len(parts)>=4 and parts[-2] in REGIONS and not ORG_RE.search(parts[-3]):
        region=parts[-2]; city=parts[-3]; rest=parts[:-3]
    elif len(parts)>=3:
        city=parts[-2]; rest=parts[:-2]
    elif len(parts)==2: rest=parts[:-1]
    else: rest=parts
    if rest:
        inst=rest[-1]; dept=', '.join(rest[:-1])
        if len(rest)>=2 and DEPT_RE.match(rest[-1]) and not DEPT_RE.match(rest[-2]):
            inst=rest[-2]; dept=', '.join(rest[:-2]+[rest[-1]])
        if len(rest)>=2 and re.search(r'Chinese Academy of Sciences|Academia Sinica|CNRS$|Helmholtz',inst,re.I) and not DEPT_RE.match(rest[-2]):
            inst=rest[-2]+', '+re.sub(r'Chinese Academy of Sciences','CAS',inst)
    return {'raw':t,'inst':inst,'dept':dept,'city':city,'region':region,'country':country}

# ---------- Institution canonicalization (light) ----------
INST_ALIASES=[
 (r'^(NHMFL|National High Magnetic Field Lab(oratory)?|National High Magnetic Field Laboratory.*|MagLab|Applied Superconductivity Center.*NHMFL.*)$','National High Magnetic Field Laboratory'),
 (r'^(Fermilab|Fermi National Accelerator Lab(oratory)?|FNAL)$','Fermi National Accelerator Laboratory'),
 (r'^(CERN|European Organization for Nuclear Research.*)$','CERN'),
 (r'^(KEK|High Energy Accelerator Research Organization.*|KEK, High Energy Accelerator Research Organization)$','KEK'),
 (r'^(NIST|National Institute of Standards and Technology.*)$','NIST'),
 (r'^(MIT|Massachusetts Institute of Technology)$','Massachusetts Institute of Technology'),
 (r'^(LBNL|Law(e)?rence Berk(e)?ley National Lab(oratory)?.*|Berkeley Lab|E\\. ?O\\. Lawrence Berkeley National Laboratory)$','Lawrence Berkeley National Laboratory'),
 (r'^(National Institute of Standards and Te\w+|NIST.*)$','NIST'),
 (r'^(University of Colorado( at)?( Boulder)?|CU Boulder)$','University of Colorado Boulder'),
 (r'^(Institute of Electrical Engineering(, CAS)?|IEE CAS|IEECAS)$','Institute of Electrical Engineering, CAS'),
 (r'^(Superconducting Magnet Division.*)$','Brookhaven National Laboratory'),
 (r'^(BNL|Brookhaven National Lab(oratory)?)$','Brookhaven National Laboratory'),
 (r'^(KIT|Karlsruhe Institute of Technology.*)$','Karlsruhe Institute of Technology'),
 (r'^(ORNL|Oak Ridge National Lab(oratory)?)$','Oak Ridge National Laboratory'),
 (r'^(PSI|Paul Scherrer Institut(e)?)$','Paul Scherrer Institute'),
 (r'^(NIFS|National Institute for Fusion Science)$','National Institute for Fusion Science'),
 (r'^(NIMS|National Institute for Materials Science)$','National Institute for Materials Science'),
 (r'^(SLAC|SLAC National Accelerator Laboratory)$','SLAC National Accelerator Laboratory'),
 (r'^(JPL|Jet Propulsion Laboratory.*)$','Jet Propulsion Laboratory'),
 (r'^(Institute of Plasma Physics, Chinese Academy of Sciences|ASIPP|Institute of Plasma Physics.*CAS.*)$','Institute of Plasma Physics, CAS'),
 (r'^(Institute of Electrical Engineering, Chinese Academy of Sciences|IEE, CAS|Institute of Electrical Engineering.*CAS.*)$','Institute of Electrical Engineering, CAS'),
 (r'^(University of Twente.*)$','University of Twente'),
 (r'^(Tohoku University.*)$','Tohoku University'),
 (r'^(Kyoto University.*)$','Kyoto University'),
 (r'^(Shanghai Jiao Tong University.*)$','Shanghai Jiao Tong University'),
 (r'^(Florida State University.*)$','Florida State University'),
 (r'^(University of Cambridge.*)$','University of Cambridge'),
]
INST_ALIASES=[(re.compile(p,re.I),c) for p,c in INST_ALIASES]
def canon_inst(name):
    n=name.strip()
    for p,c in INST_ALIASES:
        if p.match(n): return c
    return n

# ---------- Sessions & presentations ----------
SESS_RE=re.compile(r'^(\d(?:E|L|M|J|LM|PL)(?:Or|Po)?\d[A-Z]?)\s*:\s*(.*)$')
PRES_RE=re.compile(r'^(\d(?:E|L|M|J|LM|PL)(?:Or|Po)?\d[A-Z]?)-(\d+)\s*(\[[^\]]*\])?\s*:\s*(.*)$')
def parse_room(loc):
    loc=loc.replace('‚',',')
    m=re.match(r'^(.*?)\s*-\s*Level\s*(\d)\s*,\s*DLCC$',loc)
    if m: return m.group(1).strip(),int(m.group(2))
    return loc.strip(),0
byid={a['sessionid']:a for a in ag}
roots=[a for a in ag if not a['parent']]
sessions={}; events=[]
for a in roots:
    m=SESS_RE.match(a['title'])
    room,level=parse_room(a['location'])
    rec={'pid':a['sessionid'],'date':a['date'],'start':a['start'][:5],'end':a['stop'][:5],'room':room,'level':level,'cat':a['catname'],'color':a['catcolor']}
    if m:
        code=m.group(1); title=m.group(2).strip()
        mods=[]
        d=a['description'] or ''
        if 'Moderator' in d:
            txt=strip_tags(d.split('Moderator(s):')[-1])
            for line in txt.split('\n'):
                line=line.strip()
                mm=re.match(r'^(.*?)\s*\((.*)\)\s*$',line)
                if mm: mods.append({'name':mm.group(1).strip(),'org':mm.group(2).strip()})
                elif line: mods.append({'name':line,'org':''})
        rec.update({'id':code,'title':title,'chairs':mods,'items':[]})
        sessions[code]=rec
    else:
        rec.update({'title':a['title'],'id':'EV'+a['sessionid']}); events.append(rec)
# session codes -> day/area/format/slot
AREA={'E':'Electronics','L':'Large Scale','M':'Materials','J':'Joint','LM':'Large Scale / Materials','PL':'Plenary'}
DAYS={'1':'Mon','2':'Tue','3':'Wed','4':'Thu','5':'Fri'}
def decode(code):
    m=re.match(r'^(\d)(E|L|M|J|LM|PL)(Or|Po)?(\d)([A-Z])?$',code)
    return {'dayN':int(m.group(1)),'area':m.group(2),'fmt':{'Or':'Oral','Po':'Poster',None:'Plenary'}[m.group(3)],'slot':int(m.group(4)),'par':m.group(5) or ''}
for code,s in sessions.items(): s.update(decode(code))

# subfields spreadsheet
wb=openpyxl.load_workbook(W+'subfields.xlsx'); ws=wb.active
cur_sub=None; sub_of={}; type_of={}; area_of={}
for r in ws.iter_rows(values_only=True):
    if r[1]: cur_sub=re.sub(r'\s*\(\d+\)\s*$','',str(r[1])).strip()
    if r[2] and isinstance(r[2],str) and r[2].startswith('['):
        code=re.match(r'\[([^\]]+)\]',r[2]).group(1)
        sub_of[code]=cur_sub; type_of[code]=r[3]; area_of[code]=r[0]
for code,s in sessions.items():
    s['subfield']=sub_of.get(code, 'Plenary' if s['area']=='PL' else 'Other')
    s['type']=type_of.get(code, 'Plenary' if s['area']=='PL' else s['fmt'])
    if s['type'] in ('Oral','Poster'): s['type']=s['fmt']
    if '[Special Session]' in s['title'] or s['cat']=='Special Session': s['special']=True
    if s['cat']=='Memorial Session': s['memorial']=True
# moderators spreadsheet -> chairs (fills sessions without moderators in planner)
wb=openpyxl.load_workbook(W+'moderators.xlsx'); ws=wb.active
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[0] or not r[3]: continue
    code=str(r[3]).strip(); last,first=[x.strip() for x in str(r[0]).split(',',1)] if ',' in str(r[0]) else (str(r[0]),'')
    nm=(first+' '+last).strip()
    if code in sessions:
        ch=sessions[code]['chairs']
        if not any(name_key(c['name'])==name_key(nm) for c in ch): ch.append({'name':nm,'org':''})

# presentations
pres={}
media_by={m['mediaid']:m for m in md}
def parse_abstract(desc):
    i=desc.find('<hr>'); body=desc[i+4:] if i>=0 else desc
    body=re.sub(r'<script.*?</script>','',body,flags=re.S)
    j=body.rfind('</div>'); 
    if j>0: body=body[:j]
    out={'authors':[],'affs':[],'abstract':'','ack':''}
    m=re.search(r'<p><i>(.*?)</i>\s*<br />(.*?)</p>',body,re.S)
    if m:
        auth=m.group(1); afftxt=m.group(2)
        affmap={}
        for sup,txt in re.findall(r'<sup>(\d+)</sup>([^<]+)',afftxt): affmap[int(sup)]=parse_aff(txt)
        order=sorted(affmap); idx={k:i for i,k in enumerate(order)}
        out['affs']=[affmap[k] for k in order]
        for chunk in re.findall(r'([^<>,]+?)\s*(?:<sup>([\d,\s]*)</sup>)?\s*(?:,|$)',auth):
            nm,sups=chunk; nm=html.unescape(nm).strip()
            if not nm: continue
            presenting='*' in nm; nm=nm.replace('*','').strip()
            refs=[idx[int(x)] for x in re.split(r'[,\s]+',sups.strip()) if x and int(x) in idx] if sups else []
            out['authors'].append({'n':nm,'a':refs,'p':presenting})
    ma=re.search(r'<p><b>Abstract</b>\s*<br />(.*?)</p>',body,re.S)
    if ma: out['abstract']=strip_tags(ma.group(1))
    mk=re.search(r'<p><b>Acknowledgments</b>\s*<br />(.*?)</p>',body,re.S)
    if mk: out['ack']=strip_tags(mk.group(1))
    return out
kids=[a for a in ag if a['parent']]
for a in kids:
    parent=byid.get(a['parent']); 
    if not parent: continue
    pm=SESS_RE.match(parent['title'])
    if not pm: continue
    scode=pm.group(1); s=sessions[scode]
    m=PRES_RE.match(a['title'])
    room,level=parse_room(a['location'])
    rec={'pid':a['sessionid'],'session':scode,'start':a['start'][:5],'end':a['stop'][:5],'date':a['date'],'room':room,'level':level,'pos':a['position'],'speakerPid':a['speakers'] or ''}
    if m:
        rec['id']=m.group(1)+'-'+m.group(2); rec['board']=(m.group(3) or '').strip('[] '); rec['title']=m.group(4).strip()
    else:
        mm=re.match(r'^(\d(?:E|L|M|J|LM|PL)\d?)\s*:\s*(.*)$',a['title'])
        if mm and s['area']=='PL': rec['id']=mm.group(1)+'-01'; rec['title']=mm.group(2).strip()
        else: rec['id']=scode+'-x'+a['sessionid']; rec['title']=a['title'].strip(); rec['aux']=True
    rec['invited']=rec['title'].startswith('[Invited]') or s['area']=='PL'
    rec['title']=re.sub(r'^\[Invited\]\s*','',rec['title'])
    if 'Withdrawn' in rec['title']: rec['withdrawn']=True
    mrec=media_by.get(a['sessionid'])
    if mrec:
        rec.update(parse_abstract(mrec['description']))
    else: rec.update({'authors':[],'affs':[],'abstract':'','ack':''})
    pres[rec['id']]=rec; s['items'].append(rec['id'])
for s in sessions.values():
    s['items'].sort(key=lambda i:(pres[i]['start'],pres[i]['pos'],pres[i]['id']))

# ---------- Speakers table: roles ----------
spk_roles=collections.defaultdict(list)  # name_key -> list of (agenda pid, role)
spk_org={}
for s0 in sp:
    nm=(s0['firstname']+' '+s0['lastname']).strip(); k=name_key(nm)
    spk_org.setdefault(k,s0['org'])
    for pid,role in re.findall(r"urn:eventpilot:all:agenda:id:(\d+)'>.*?</a>\s*-\s*(\w+)",s0['description'] or ''):
        spk_roles[k].append((pid,role))
pid2pres={p['pid']:p['id'] for p in pres.values()}; pid2sess={s['pid']:s['id'] for s in sessions.values()}

# ---------- People ----------
# cluster by name key; split clusters when affiliations are disjoint
entries=collections.defaultdict(list)  # key -> list of (pres id, author idx, inst set, display name)
for p in pres.values():
    for ai,au in enumerate(p['authors']):
        k=name_key(au['n'])
        insts=set(norm(canon_inst(p['affs'][x]['inst'])) for x in au['a'] if x<len(p['affs']))
        ctry=set(p['affs'][x]['country'] for x in au['a'] if x<len(p['affs']))
        entries[k].append((p['id'],ai,insts,au['n'],ctry))
people=[]; 
for k,lst in entries.items():
    # union-find on shared institution
    n=len(lst); parent=list(range(n))
    def find(x):
        while parent[x]!=x: parent[x]=parent[parent[x]]; x=parent[x]
        return x
    for i in range(n):
        for j in range(i+1,n):
            if lst[i][2]&lst[j][2] or lst[i][4]&lst[j][4] or not lst[i][2] or not lst[j][2]: parent[find(i)]=find(j)
    groups=collections.defaultdict(list)
    for i in range(n): groups[find(i)].append(i)
    # merge groups that would be identical-named singletons? keep separate but label
    for gi,idxs in enumerate(groups.values()):
        names=collections.Counter(lst[i][3] for i in idxs)
        disp=max(names.items(),key=lambda x:(len(x[0]),x[1]))[0]
        person={'id':k.replace(' ','_')+('' if len(groups)==1 else '_%d'%gi),'name':disp,'key':k,'pres':[],'insts':collections.Counter(),'countries':collections.Counter(),'presenting':0,'invited':0}
        for i in idxs:
            pidx,ai,insts,_,_c=lst[i]; p=pres[pidx]; au=p['authors'][ai]
            person['pres'].append(pidx)
            au['pid']=person['id']
            if au['p']: person['presenting']+=1
            if au['p'] and p.get('invited'): person['invited']+=1
            for x in au['a']:
                if x<len(p['affs']):
                    person['insts'][canon_inst(p['affs'][x]['inst'])]+=1; person['countries'][p['affs'][x]['country']]+=1
        people.append(person)
pk={}
for person in people: pk.setdefault(person['key'],[]).append(person)
# chairs -> person ids (match by key; if multiple, prefer org overlap)
def match_person(nm,org=''):
    k=name_key(nm); cands=pk.get(k)
    if not cands: return None
    if len(cands)==1: return cands[0]
    if org:
        o=norm(org)
        for c in cands:
            if any(norm(i) in o or o in norm(i) for i in c['insts']): return c
    return max(cands,key=lambda c:len(c['pres']))
for s in sessions.values():
    for ch in s['chairs']:
        per=match_person(ch['name'],ch['org'])
        if per: ch['pid']=per['id']; per.setdefault('chairs',[]).append(s['id'])
        else:
            # create a person record for chairs that don't author anything
            k=name_key(ch['name']); per={'id':'chair_'+k.replace(' ','_'),'name':ch['name'],'key':k,'pres':[],'insts':collections.Counter({ch['org']:1} if ch['org'] else {}),'countries':collections.Counter(),'presenting':0,'invited':0,'chairs':[s['id']]}
            people.append(per); pk.setdefault(k,[]).append(per); ch['pid']=per['id']
# speaker table roles (Moderator for sessions not captured, org fallback)
for k,lst in spk_roles.items():
    cands=pk.get(k)
    if not cands: continue
    per=max(cands,key=lambda c:len(c['pres']))
    for pid,role in lst:
        if role=='Moderator' and pid in pid2sess:
            sc=pid2sess[pid]
            if sc not in per.setdefault('chairs',[]): per['chairs'].append(sc)
            s=sessions[sc]
            if not any(c.get('pid')==per['id'] for c in s['chairs']): s['chairs'].append({'name':per['name'],'org':spk_org.get(k,''),'pid':per['id']})
    if not per['insts'] and spk_org.get(k): per['insts'][spk_org[k]]=1

# ---------- Plenary speakers (bios + photos) ----------
def img_b64(path,maxw=200):
    try:
        from PIL import Image
        im=Image.open(path).convert('RGB'); w,h=im.size
        tw,th=maxw,int(maxw*4/3)
        r=max(tw/w,th/h); im=im.resize((max(1,int(w*r)),max(1,int(h*r))))
        w2,h2=im.size; left=(w2-tw)//2; top=max(0,(h2-th)//3)
        im=im.crop((left,top,left+tw,top+th))
        b=io.BytesIO(); im.save(b,'JPEG',quality=72,optimize=True); return 'data:image/jpeg;base64,'+base64.b64encode(b.getvalue()).decode()
    except Exception as e:
        print('img fail',path,e,file=sys.stderr); return ''
PL_PRES=[p for p in pres.values() if sessions[p['session']]['area']=='PL' and p['authors']]
def pl_match(lastname):
    ln=norm(lastname)
    for p in PL_PRES:
        for au in p['authors']:
            if au['p'] and norm(au['n']).split()[-1]==ln: return p,au
    return None,None
PHOTO_LAST={'paul-chu':'Chu','irfan-siddiqi':'Sidiqqi','john-clarke':'Clarke','sam-benz':'Benz','rossilucio':'Rossi','timothy-murphy':'Murphy','max-propper':'Pröpper','roberta-satariano':'Satariano','raphael-unterrainer':'Unterrainer','jeseok-bang':'Bang','nicolo-riva':'Riva','shaon-barua':'Barua','tugrul-ersoz':'Ersoz'}
t=open(W+'plenary.html',encoding='utf-8',errors='ignore').read()
body=re.sub(r'<(script|style|header|nav|footer).*?</\1>','',t,flags=re.S)
plen=[]
blocks=re.split(r'(?=<img[^>]+src="[^"]+/uploads/sites/6/2026/\d\d/[^"]+\.(?:jpg|jpeg|png)")',body)
for b in blocks[1:]:
    img=re.search(r'src="([^"]+)"',b).group(1); fn=img.split('/')[-1]
    stem=[k for k in PHOTO_LAST if fn.startswith(k)]
    if not stem: continue
    p,au=pl_match(PHOTO_LAST[stem[0]])
    names=[strip_tags(x) for x in re.findall(r'<strong>(.*?)</strong>',b[:4000])]
    names=[x for x in names if ':' not in x and 2<=len(x.split())<=5 and norm(PHOTO_LAST[stem[0]]).replace('sidiqqi','siddiqi') in norm(x)]
    name=names[0] if names else (au['n'] if au else stem[0])
    if name=='Irfan Siddiqi' and au: au['n']='Irfan Siddiqi'
    titles=[strip_tags(x) for x in re.findall(r'<em>(.*?)</em>',b[:3000])][:3]
    paras=[strip_tags(x) for x in re.findall(r'<p[^>]*>(.*?)</p>',b[:12000],re.S)]
    bio=' '.join(x for x in paras if len(x)>60 and not x.startswith(('Presentation','Abstract')) and x not in titles)
    pt=re.search(r'Presentation title:</strong>\s*(.*?)</p>',b,re.S)
    rec={'name':name,'titles':titles,'bio':bio.strip()[:3000],'talk':strip_tags(pt.group(1)) if pt else (p['title'] if p else ''),'photo':img_b64(W+'photos/'+fn),'photoSrc':img,'pres':p['id'] if p else ''}
    plen.append(rec)
# plenary abstracts from the ASC website
for a,last in [('benz','Benz'),('siddiqi','Siddiqi'),('siddiqi','Sidiqqi'),('rossi','Rossi'),('murphy','Murphy'),('clarke','Clarke'),('chu','Chu')]:
    p,au=pl_match(last)
    if not p or p['abstract']: continue
    tt=open(W+'plen_%s.html'%a,encoding='utf-8',errors='ignore').read()
    bb=re.sub(r'<(script|style|header|nav|footer).*?</\1>','',tt,flags=re.S)
    ps=[strip_tags(x) for x in re.findall(r'<p[^>]*>(.*?)</p>',bb,re.S)]
    ps=[x for x in ps if len(x)>150 and not x.startswith(('Copyright','To receive'))]
    p['abstract']='\n'.join(ps); p['abstractSrc']='https://www.appliedsuperconductivity.org/asc2026/%s-abstract/'%a
for pl in plen:
    per=None
    if pl['pres']:
        p=pres[pl['pres']]
        for au in p['authors']:
            if au['p'] and au.get('pid'): per=next((x for x in people if x['id']==au['pid']),None); break
    if not per: per=match_person(pl['name'])
    if per:
        per['plenary']=True; per['bio']=pl['bio']; per['titles']=pl['titles']; per['photo']=pl['photo']; per['photoCredit']='appliedsuperconductivity.org'; pl['pid']=per['id']
        # the conference website spelling is authoritative for plenary speakers (the planner has typos)
        if len(pl['name'].split())>=2:
            per['name']=pl['name']
            for pid_ in per['pres']:
                for au in pres[pid_]['authors']:
                    if au.get('pid')==per['id']: au['n']=pl['name']
    else: print('plenary unmatched',pl['name'],file=sys.stderr)

# ---------- Special sessions descriptions ----------
t=open(W+'special-sessions.html',encoding='utf-8',errors='ignore').read()
body=re.sub(r'<(script|style|header|nav|footer).*?</\1>','',t,flags=re.S)
specials=[]
for m in re.finditer(r'<(h[2-5]|strong)[^>]*>(.*?)</\1>\s*(?:</p>)?\s*((?:<p[^>]*>.*?</p>\s*){1,4})',body,re.S):
    title=strip_tags(m.group(2)).strip(); desc=strip_tags(m.group(3)).strip()
    if 20<len(title)<160 and len(desc)>120: specials.append({'title':title,'desc':desc})
# match to sessions by title similarity
def simkey(s): return set(w for w in norm(s).split() if len(w)>3)
for sps in specials:
    best=None;bs=0
    for s in sessions.values():
        if not (s.get('special') or s.get('memorial')): continue
        a=simkey(sps['title']); b=simkey(s['title']); sc=len(a&b)/max(1,len(a|b))
        if sc>bs: bs=sc;best=s
    if best and bs>0.4: best['desc']=sps['desc']

# ---------- Output ----------
people.sort(key=lambda p:(-len(p['pres']),p['name']))
for per in people:
    per['insts']=[i for i,_ in per['insts'].most_common()]; per['countries']=[c for c,_ in per['countries'].most_common()]
    per['pres'].sort()
inst_rows=collections.Counter(); country_rows=collections.Counter(); inst_country={}
for p in pres.values():
    seen=set()
    for a in p['affs']:
        ci=canon_inst(a['inst'])
        if ci not in seen: inst_rows[ci]+=1; seen.add(ci)
        inst_country.setdefault(ci,collections.Counter())[a['country']]+=1
    for c in set(a['country'] for a in p['affs']): country_rows[c]+=1
insts=[{'name':n,'count':c,'country':inst_country[n].most_common(1)[0][0]} for n,c in inst_rows.most_common()]
subfields=collections.defaultdict(list)
for s in sessions.values(): subfields[(s['area'],s['subfield'])].append(s['id'])
data={
 'meta':{'name':'ASC 2026','venue':'David L. Lawrence Convention Center, Pittsburgh','tz':'America/New_York','built':datetime.datetime.utcnow().isoformat()+'Z','dates':['2026-09-06','2026-09-07','2026-09-08','2026-09-09','2026-09-10','2026-09-11'],'areas':AREA,'source':'https://eppro01.ativ.me/web/planner.php?id=ASC2026'},
 'sessions':sorted(sessions.values(),key=lambda s:(s['date'],s['start'],s['id'])),
 'presentations':sorted(pres.values(),key=lambda p:(p['date'],p['start'],p['id'])),
 'people':people,'events':sorted(events,key=lambda e:(e['date'],e['start'])),
 'institutions':insts,'countries':[{'name':c,'count':n} for c,n in country_rows.most_common()],
 'subfields':[{'area':a,'name':n,'sessions':sorted(v)} for (a,n),v in sorted(subfields.items())],
 'plenary':plen,
}
for p in data['presentations']:
    for a in p['affs']: a['inst']=canon_inst(a['inst'])
json.dump(data,open(S+'/build/data.json','w'),ensure_ascii=False)
print('sessions',len(sessions),'presentations',len(pres),'with abstract',sum(1 for p in pres.values() if p['abstract']),'people',len(people),'insts',len(insts),'countries',len(country_rows),'events',len(events),'plenary',len(plen),'specials matched',sum(1 for s in sessions.values() if s.get('desc')))
print('size',os.path.getsize(S+'/build/data.json'))
